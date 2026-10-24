import html
import logging
import re
import time
import zipfile
from pathlib import Path

import requests
from django.conf import settings
from django.core.cache import cache

logger = logging.getLogger(__name__)

GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
)
REQUEST_TIMEOUT = 45
MAX_QUESTION_CHARS = 2000
MAX_HISTORY_MESSAGES = 10
PER_FILE_CHAR_LIMIT = 20_000
TOTAL_CONTEXT_CHAR_LIMIT = 120_000
MAX_PDF_PAGES = 150
THROTTLE_LIMIT = 15
THROTTLE_WINDOW_SECONDS = 60
FALLBACK_MODELS = ["gemini-3.7-flash", "gemini-3.5-flash", "gemini-flash-lite-latest"]
OVERLOAD_STATUSES = {404, 429, 500, 503}
MODEL_GONE_MARKERS = (
    "no longer available",
    "is not supported",
    "not found",
    "does not exist",
    "not supported for",
)
ATTEMPT_PAUSE_SECONDS = 0.9

# Model that last answered successfully; tried first on subsequent calls.
_preferred_model = None


class AiNotConfigured(Exception):
    """AI can't be used because of setup/credentials. str(exc) is user-facing."""


class AiUpstreamError(Exception):
    pass


SYSTEM_PROMPT = """You are StudyHub's study assistant, embedded next to a student's course material.

Rules:
- Prefer answering from the STUDY MATERIAL provided below when it is relevant.
- If the material does not cover the question, answer from your own knowledge and say the notes did not cover it.
- Be clear, concise and exam-oriented. Use short paragraphs, bullet points and simple examples.
- Use Markdown formatting (**bold**, bullet lists) sparingly.
- Never invent citations or pretend parts of the material exist when they do not.
"""

SLIDE_NAME_RE = re.compile(r"ppt/slides/slide(\d+)\.xml")
TEXT_TAG_RE = re.compile(r"<a:t>(.*?)</a:t>", re.S)
ODP_PARA_RE = re.compile(r"<text:[ph][^>]*>(.*?)</text:[ph]>", re.S)


def _clean_xml_text(raw):
    text = re.sub(r"<[^>]+>", " ", raw)
    return html.unescape(text).strip()


def _extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks = []
    for i, page in enumerate(reader.pages):
        if i >= MAX_PDF_PAGES:
            break
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(chunks)


def _extract_txt(path):
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def _extract_docx(path):
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_pptx(path):
    texts = []
    with zipfile.ZipFile(path) as z:
        slides = []
        for name in z.namelist():
            match = SLIDE_NAME_RE.fullmatch(name)
            if match:
                slides.append((int(match.group(1)), name))
        for _, name in sorted(slides):
            xml = z.read(name).decode("utf-8", "ignore")
            found = TEXT_TAG_RE.findall(xml)
            texts.append("\n".join(html.unescape(t) for t in found))
    return "\n\n".join(texts)


def _extract_odp(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("content.xml").decode("utf-8", "ignore")
    paras = [_clean_xml_text(p) for p in ODP_PARA_RE.findall(xml)]
    return "\n".join(p for p in paras if p)


def extract_file_text(file_field):
    """Extract plain text from a Material's file (or its converted PDF)."""
    field = file_field.converted_pdf or file_field.file
    try:
        path = field.path
    except Exception:
        return ""
    ext = Path(field.name).suffix.lower()
    extractors = {
        ".pdf": _extract_pdf,
        ".txt": _extract_txt,
        ".docx": _extract_docx,
        ".pptx": _extract_pptx,
        ".ppsx": _extract_pptx,
        ".odp": _extract_odp,
    }
    extractor = extractors.get(ext)
    if not extractor:
        return ""
    try:
        return extractor(path)[:PER_FILE_CHAR_LIMIT]
    except Exception:
        logger.exception("AI text extraction failed for %s", field.name)
        return ""


def build_context(topics):
    """Build a single text blob describing all materials under the given topics."""
    sections = []
    total = 0
    for topic in topics:
        for material in topic.materials.all():
            body = extract_file_text(material)
            section = f"[File: {material.title} ({material.get_file_type_display()})]\n{body}".strip()
            remaining = TOTAL_CONTEXT_CHAR_LIMIT - total
            if remaining <= 0:
                return "\n\n---\n\n".join(sections)
            sections.append(section[:remaining])
            total += min(len(section), remaining)
    return "\n\n---\n\n".join(sections)


def throttled(user_id):
    key = f"ai_throttle_{user_id}"
    hits = cache.get(key)
    if hits is None:
        cache.set(key, 1, THROTTLE_WINDOW_SECONDS)
        return False
    if hits >= THROTTLE_LIMIT:
        return True
    cache.set(key, hits + 1, THROTTLE_WINDOW_SECONDS)
    return False


def _parse_answer(data):
    block_reason = (data.get("promptFeedback") or {}).get("blockReason")
    if block_reason:
        raise AiUpstreamError(f"The request was blocked ({block_reason}). Try rephrasing.")
    candidates = data.get("candidates") or []
    if not candidates:
        raise AiUpstreamError("The AI returned no answer. Try again.")
    parts = ((candidates[0].get("content") or {}).get("parts")) or []
    answer = "".join(part.get("text", "") for part in parts).strip()
    if not answer:
        finish = candidates[0].get("finishReason")
        if finish == "SAFETY":
            raise AiUpstreamError("The answer was blocked by safety filters. Try rephrasing.")
        raise AiUpstreamError("The AI returned an empty answer. Try again.")
    return answer


def ask_gemini(question, history=None, context=""):
    global _preferred_model
    api_key = settings.GEMINI_API_KEY
    if not api_key:
        raise AiNotConfigured()

    system_text = SYSTEM_PROMPT
    if context:
        system_text += (
            "\n\n=== STUDY MATERIAL START ===\n"
            f"{context}\n"
            "=== STUDY MATERIAL END ==="
        )

    contents = []
    for msg in (history or [])[-MAX_HISTORY_MESSAGES:]:
        role = "model" if msg.get("role") == "model" else "user"
        text = str(msg.get("text") or "")[:MAX_QUESTION_CHARS]
        if text:
            contents.append({"role": role, "parts": [{"text": text}]})
    contents.append({"role": "user", "parts": [{"text": question}]})

    payload = {
        "system_instruction": {"parts": [{"text": system_text}]},
        "contents": contents,
        "generationConfig": {"temperature": 0.4, "maxOutputTokens": 1024},
    }

    models = [settings.GEMINI_MODEL] if settings.GEMINI_MODEL else []
    for model in FALLBACK_MODELS:
        if model not in models:
            models.append(model)
    if _preferred_model and _preferred_model in models:
        models.remove(_preferred_model)
        models.insert(0, _preferred_model)

    last_overload = None
    for index, model in enumerate(models):
        started = time.monotonic()
        try:
            response = requests.post(
                GEMINI_URL.format(model=model),
                headers={"x-goog-api-key": api_key},
                json=payload,
                timeout=REQUEST_TIMEOUT,
            )
        except requests.RequestException as exc:
            logger.exception("Gemini request failed")
            raise AiUpstreamError(
                f"Could not reach the AI service ({exc.__class__.__name__})."
            )
        logger.info(
            "Gemini model %s took %.1fs -> %s",
            model,
            time.monotonic() - started,
            response.status_code,
        )

        if response.status_code == 200:
            _preferred_model = model
            return _parse_answer(response.json())

        if response.status_code in (401, 403):
            raise AiNotConfigured(
                f"The AI service rejected the API key ({response.status_code}). "
                "Double-check GEMINI_API_KEY in the .env file."
            )
        if response.status_code == 400 and (
            "API key not valid" in response.text
            or "API_KEY_INVALID" in response.text
        ):
            raise AiNotConfigured(
                "The AI service says the API key is not valid. Replace GEMINI_API_KEY in "
                "the .env file with a valid key from aistudio.google.com/apikey."
            )

        if response.status_code == 400 and any(
            marker in response.text.lower() for marker in MODEL_GONE_MARKERS
        ):
            last_overload = response.status_code
            logger.warning("Model %s unavailable, trying next fallback", model)
            if index < len(models) - 1:
                time.sleep(ATTEMPT_PAUSE_SECONDS)
            continue

        if response.status_code in OVERLOAD_STATUSES:
            last_overload = response.status_code
            logger.warning(
                "Model %s busy/unavailable (%s), trying next fallback",
                model,
                response.status_code,
            )
            if index < len(models) - 1:
                time.sleep(ATTEMPT_PAUSE_SECONDS)
            continue

        logger.error("Gemini error %s: %s", response.status_code, response.text[:300])
        raise AiUpstreamError(
            f"The AI service returned an error ({response.status_code})."
        )

    raise AiUpstreamError(
        "All AI models are busy right now (free-tier capacity). Please try again in a few seconds."
    )